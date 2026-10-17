import json
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from dendrobase.config import Settings
from dendrobase.excel_import import (STATE_FILE, ImportProgress, ImportStats, excel_files, fill_journal_numbers,
                                     import_folder, map_param, probe_name_from_file, read_crn_sheet,
                                     read_series_sheet, read_title_sheet, run)
from dendrobase.models import CrnRow, SeriesValue


def write_workbook(path: Path, *, title=None, crn=None, series=None) -> Path:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    if title is not None:
        sheet = workbook.create_sheet("Title")
        for row in title:
            sheet.append(row)
    if crn is not None:
        sheet = workbook.create_sheet("crn")
        for row in crn:
            sheet.append(row)
    if series is not None:
        sheet = workbook.create_sheet("series")
        for row in series:
            sheet.append(row)
    workbook.save(path)
    return path


@pytest.mark.parametrize("param, column", [
    ("Широта,грд", "shirota_grd"),
    ("  Главная   порода ", "glavnaya_poroda"),
    ("Коэфф. состава", "koef_sostava"),          # не блытаецца з «состав»
    ("Состав древостоя", "sostav_drevoj"),
    ("Геоботанический район (по карте)", "geobotanicheskij_rajon"),  # не «район»
    ("Район", "rayon"),
    ("N журнала", "n_journal"),
    ("Нешта невядомае", None),
    ("", None),
])
def test_map_param(param, column):
    assert map_param(param) == column


def test_probe_name_from_file():
    assert probe_name_from_file(Path("/x/PP_01.xlsx")) == "pp_01"


def test_read_title_sheet(tmp_path):
    path = write_workbook(tmp_path / "PP01.xlsx", title=[
        ("Широта,грд", 54.1), ("Долгота, грд", "26,2"), ("Главная порода", "Е"),
        ("Первый год", 1850.0), ("Дата отбора образцов", "15.08.2019"), ("Полнота", "0,8"),
        ("Коэфф. состава", 7), ("Состав", "7Е3С"), ("Квартал", 12.0), ("Возраст средний", "80-120"),
        ("N пробной площади", "іншая-назва"), (None, "пуста"), ("Площадь выдела", None),
    ])
    data, errors = read_title_sheet(path, "pp01")
    assert data["n_probnoy_ploshchadi"] == "pp01"  # з назвы файла, не з ліста
    assert data["shirota_grd"] == 54.1 and data["dolgota_grd"] == 26.2
    assert data["pervyy_god"] == 1850
    assert data["data_otbora_obraztsov"] == date(2019, 8, 15)
    assert data["polnota"] == 0.8
    assert data["koef_sostava"] == 7.0 and data["sostav_drevoj"] == "7Е3С"
    assert data["kvartal"] == "12"
    assert "vozrast" not in data and errors == ["Сярэдні ўзрост: '80-120'"]  # подпіс поля беларускі, параметр Excel рускі
    assert "ploshchad_vydela" not in data


def test_read_crn_sheet(tmp_path):
    path = write_workbook(tmp_path / "pp.xlsx", crn=[
        ("Year", "num", "RAW", "std", "res", "ars", "lishniy"),
        (1850, 10, 1.1, 1.0, 0.99, 1.01, 5),
        (1851, None, "1,2", None, None, None, None),
        (1851, 1, 9, 9, 9, 9, 9),
        (None, 1, 1, 1, 1, 1, 1),
        ("abc", 1, 1, 1, 1, 1, 1),
    ])
    rows, errors = read_crn_sheet(path, "pp")
    assert rows == [CrnRow(1850, 10.0, 1.1, 1.0, 0.99, 1.01), CrnRow(1851, None, 1.2, None, None, None)]
    assert len(errors) == 2


def test_read_series_sheet(tmp_path):
    path = write_workbook(tmp_path / "pp.xlsx", series=[
        ("year", 1, "1a", "2"), (1850, 1.1, 1.2, "1,3"), (1851, 1.4, None, "x"), (None, 1, 1, 1),
    ])
    values, errors = read_series_sheet(path, "pp")
    assert values == [
        SeriesValue(1850, "1", 1.1), SeriesValue(1850, "1a", 1.2), SeriesValue(1850, "2", 1.3),
        SeriesValue(1851, "1", 1.4),
    ]
    assert len(errors) == 1 and "серыя 2" in errors[0]


class RecordingDb:
    """Дубль DatabaseManager: запамінае выклікі, нічога не піша."""

    def __init__(self, fail_for: str | None = None):
        self.calls: list[tuple] = []
        self.probes: set[str] = set()
        self.fail_for = fail_for

    def get_all_probes(self):
        return sorted(self.probes)

    def upsert_title(self, data):
        from dendrobase.db import DatabaseError
        if data["n_probnoy_ploshchadi"] == self.fail_for:
            raise DatabaseError("штучная памылка")
        self.calls.append(("title", dict(data)))
        self.probes.add(data["n_probnoy_ploshchadi"])

    def upsert_crn_rows(self, probe, rows, *, overwrite=True):
        rows = list(rows)
        self.calls.append(("crn", probe, rows, overwrite))
        return len(rows)

    def upsert_series_values(self, probe, values, *, overwrite=True):
        values = list(values)
        self.calls.append(("series", probe, values, overwrite))
        return len(values)

    def set_journal_number(self, probe, value):
        self.calls.append(("journal", probe, value))
        return 1


def test_import_folder_processes_each_file_and_counts(tmp_path):
    write_workbook(tmp_path / "PP01.xlsx", title=[("Широта,грд", 54.1)],
                   crn=[("year", "raw"), (1850, 1.1)], series=[("year", "1"), (1850, 1.5)])
    write_workbook(tmp_path / "PP02.xlsx", title=[("Главная порода", "С")])
    write_workbook(tmp_path / "no_title.xlsx", crn=[("year", "raw"), (1850, 1.1)])
    write_workbook(tmp_path / "~$PP01.xlsx", title=[("Широта,грд", 1)])
    write_workbook(tmp_path / "dendro_point_living.xlsx", title=[("Широта,грд", 1)])
    (tmp_path / "notes.txt").write_text("x")

    assert [p.name for p in excel_files(tmp_path)] == ["PP01.xlsx", "PP02.xlsx", "no_title.xlsx"]

    db = RecordingDb()
    stats = import_folder(db, tmp_path)
    assert (stats.files, stats.title_updated, stats.crn_rows, stats.series_rows, stats.warnings, stats.errors) == (3, 2, 1, 1, 1, 0)
    kinds = [call[0] for call in db.calls]
    assert kinds == ["title", "crn", "series", "title"]
    assert db.calls[1][3] is False  # па змаўчанні існуючыя радкі не перапісваюцца
    assert db.calls[0][1]["n_probnoy_ploshchadi"] == "pp01"


def test_import_folder_continues_after_database_error(tmp_path):
    write_workbook(tmp_path / "bad.xlsx", title=[("Широта,грд", 1)])
    write_workbook(tmp_path / "good.xlsx", title=[("Широта,грд", 2)])
    db = RecordingDb(fail_for="bad")
    stats = import_folder(db, tmp_path)
    assert stats.errors == 1 and stats.title_updated == 1
    assert "памылак: 1" in stats.summary()


def test_fill_journal_numbers(tmp_path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Название пробной площади", "N журнала"])
    sheet.append(["PP01", 42])
    sheet.append(["PP02", 0])
    sheet.append([None, 7])
    path = tmp_path / "dendro_point_living.xlsx"
    workbook.save(path)
    db = RecordingDb()
    assert fill_journal_numbers(db, path) == 1
    assert db.calls == [("journal", "PP01", "42")]


def test_fill_journal_numbers_requires_columns(tmp_path):
    workbook = openpyxl.Workbook()
    workbook.active.append(["a", "b"])
    path = tmp_path / "dendro_point_living.xlsx"
    workbook.save(path)
    with pytest.raises(ValueError):
        fill_journal_numbers(RecordingDb(), path)


def test_import_stats_summary_mentions_all_counters():
    text = ImportStats(files=3, skipped=4, title_updated=2, crn_rows=5, series_rows=7, warnings=1, errors=0).summary()
    assert "Файлаў: 3" in text and "прапушчана ўжо апрацаваных: 4" in text and "SERIES: 7" in text


# ------------------------------------------------------------- стан імпарту

def read_state(folder: Path) -> dict:
    return json.loads((folder / STATE_FILE).read_text(encoding="utf-8"))


def test_second_run_skips_unchanged_files(tmp_path):
    write_workbook(tmp_path / "PP01.xlsx", title=[("Широта,грд", 54.1)], crn=[("year", "raw"), (1850, 1.1)])
    write_workbook(tmp_path / "PP02.xlsx", title=[("Главная порода", "С")])
    db = RecordingDb()
    first = import_folder(db, tmp_path)
    assert (first.files, first.skipped) == (2, 0)

    state = read_state(tmp_path)
    assert state["format"] == 1 and list(state["files"]) == ["PP01.xlsx", "PP02.xlsx"]
    entry = state["files"]["PP01.xlsx"]
    assert entry["status"] == "done" and entry["probe"] == "pp01" and entry["overwrite"] is False
    assert len(entry["sha256"]) == 64 and entry["warnings"] == [] and "error" not in entry

    db.calls.clear()
    second = import_folder(db, tmp_path)
    assert (second.files, second.skipped, second.title_updated) == (0, 2, 0)
    assert db.calls == []
    assert "прапушчана ўжо апрацаваных: 2" in second.summary()


def test_changed_file_is_imported_again(tmp_path):
    path = write_workbook(tmp_path / "PP01.xlsx", title=[("Широта,грд", 54.1)])
    write_workbook(tmp_path / "PP02.xlsx", title=[("Широта,грд", 1)])
    db = RecordingDb()
    import_folder(db, tmp_path)
    write_workbook(path, title=[("Широта,грд", 55.0)])

    stats = import_folder(db, tmp_path)
    assert (stats.files, stats.skipped) == (1, 1)
    assert db.calls[-1][1] == {"n_probnoy_ploshchadi": "pp01", "shirota_grd": 55.0}


def test_state_keeps_warnings_and_errors_and_retries_failed_files(tmp_path):
    write_workbook(tmp_path / "PP01.xlsx", title=[("Широта,грд", 54.1), ("Возраст средний", "80-120")],
                   crn=[("year", "raw"), ("abc", 1)])
    write_workbook(tmp_path / "bad.xlsx", title=[("Широта,грд", 1)])
    write_workbook(tmp_path / "no_title.xlsx", crn=[("year", "raw"), (1850, 1.1)])
    db = RecordingDb(fail_for="bad")
    import_folder(db, tmp_path)

    files = read_state(tmp_path)["files"]
    assert files["PP01.xlsx"]["status"] == "done"
    assert files["PP01.xlsx"]["warnings"][0] == "title, поле прапушчана: Сярэдні ўзрост: '80-120'"
    assert files["PP01.xlsx"]["warnings"][1].startswith("crn, радок 2")
    assert files["bad.xlsx"]["status"] == "error" and "штучная памылка" in files["bad.xlsx"]["error"]
    assert files["no_title.xlsx"]["status"] == "skipped"
    assert files["no_title.xlsx"]["warnings"] == ["няма ліста title, файл прапушчаны"]

    db.fail_for = None  # памылковы файл спрабуецца зноў; з папярэджаннямі і без title — прапускаюцца
    stats = import_folder(db, tmp_path)
    assert (stats.files, stats.skipped, stats.errors, stats.title_updated) == (1, 2, 0, 1)
    assert read_state(tmp_path)["files"]["bad.xlsx"]["status"] == "done"


def test_overwrite_run_reimports_files_imported_without_it(tmp_path):
    write_workbook(tmp_path / "PP01.xlsx", title=[("Широта,грд", 54.1)], crn=[("year", "raw"), (1850, 1.1)])
    db = RecordingDb()
    import_folder(db, tmp_path)

    stats = import_folder(db, tmp_path, overwrite=True)
    assert (stats.files, stats.skipped) == (1, 0) and db.calls[-1][3] is True
    assert read_state(tmp_path)["files"]["PP01.xlsx"]["overwrite"] is True
    assert import_folder(db, tmp_path, overwrite=True).skipped == 1
    assert import_folder(db, tmp_path).skipped == 1  # звычайны запуск пасля --overwrite таксама прапускае


def test_reimport_flag_and_missing_probe_ignore_state(tmp_path):
    write_workbook(tmp_path / "PP01.xlsx", title=[("Широта,грд", 54.1)])
    db = RecordingDb()
    import_folder(db, tmp_path)
    assert import_folder(db, tmp_path, reimport=True).files == 1
    assert import_folder(db, tmp_path).skipped == 1

    db.probes.clear()  # базу ачысцілі: ПП больш няма, стан не лічыцца
    assert import_folder(db, tmp_path).files == 1


def test_corrupt_state_file_is_ignored(tmp_path, caplog):
    write_workbook(tmp_path / "PP01.xlsx", title=[("Широта,грд", 54.1)])
    (tmp_path / STATE_FILE).write_text("{не json", encoding="utf-8")
    stats = import_folder(RecordingDb(), tmp_path)
    assert stats.files == 1 and "не прачытаны" in caplog.text
    assert read_state(tmp_path)["files"]["PP01.xlsx"]["status"] == "done"


def test_unwritable_state_does_not_stop_import(tmp_path, caplog):
    write_workbook(tmp_path / "PP01.xlsx", title=[("Широта,грд", 54.1)])
    write_workbook(tmp_path / "PP02.xlsx", title=[("Широта,грд", 1)])
    (tmp_path / STATE_FILE).mkdir()  # запісаць файл на месца папкі нельга
    stats = import_folder(RecordingDb(), tmp_path)
    assert (stats.files, stats.errors) == (2, 0)
    assert caplog.text.count("Стан імпарту не захаваны") == 1
    assert not (tmp_path / (STATE_FILE + ".tmp")).exists()


# ------------------------------------------------------ ход працы і спыненне (для акна імпарту)

def test_progress_reports_scan_each_file_and_unchanged(tmp_path):
    write_workbook(tmp_path / "PP01.xlsx", title=[("Широта,грд", 54.1)],
                   crn=[("year", "raw"), (1850, 1.1), (1851, 1.2)], series=[("year", "1"), (1850, 1.5)])
    write_workbook(tmp_path / "no_title.xlsx", crn=[("year", "raw"), (1850, 1.1)])
    db = RecordingDb()
    events: list[ImportProgress] = []
    import_folder(db, tmp_path, progress=events.append)
    assert [(e.stage, e.done, e.total, e.file) for e in events] == [
        ("scan", 0, 2, ""), ("file", 1, 2, "PP01.xlsx"), ("file", 2, 2, "no_title.xlsx")]
    first = events[1].result
    assert first.status == "done" and first.probe == "pp01" and (first.crn_rows, first.series_rows) == (2, 1)
    assert events[2].result.status == "skipped" and events[2].result.warnings == ["няма ліста title, файл прапушчаны"]

    events.clear()
    import_folder(db, tmp_path, progress=events.append)
    assert [(e.stage, e.file, e.result) for e in events[1:]] == [
        ("unchanged", "PP01.xlsx", None), ("unchanged", "no_title.xlsx", None)]


def test_unreadable_file_is_reported_as_error_event(tmp_path, monkeypatch):
    write_workbook(tmp_path / "PP01.xlsx", title=[("Широта,грд", 54.1)])
    monkeypatch.setattr("dendrobase.excel_import.file_digest", lambda path: (_ for _ in ()).throw(OSError("няма доступу")))
    events: list[ImportProgress] = []
    stats = import_folder(RecordingDb(), tmp_path, progress=events.append)
    assert (stats.files, stats.errors) == (1, 1)
    assert events[1].result.status == "error" and "няма доступу" in events[1].result.error


def test_should_stop_cancels_between_files_and_keeps_state(tmp_path):
    for name in ("a.xlsx", "b.xlsx", "c.xlsx"):
        write_workbook(tmp_path / name, title=[("Широта,грд", 1)])
    db = RecordingDb()
    answers = iter([False, True])  # перад першым файлам «працуй», перад другім «спыніся»
    stats = import_folder(db, tmp_path, should_stop=lambda: next(answers))
    assert stats.cancelled and (stats.files, stats.errors) == (1, 0)
    assert [call[1]["n_probnoy_ploshchadi"] for call in db.calls] == ["a"]
    assert list(read_state(tmp_path)["files"]) == ["a.xlsx"]  # зробленае запісана ў стан

    rest = import_folder(db, tmp_path)  # паўторны запуск дабірае астатнія і не чапае зробленае
    assert (rest.files, rest.skipped, rest.cancelled) == (2, 1, False)


def test_run_imports_into_database_and_reports_stages(clean_db, tmp_path):
    """Поўны імпарт праз run(), як у акне імпарту: уласнае злучэнне, этапы journal і geometry, зводка."""
    write_workbook(tmp_path / "PP01.xlsx", title=[("Широта,грд", 54.1), ("Долгота, грд", 26.2), ("Главная порода", "С")],
                   crn=[("year", "raw"), (1850, 1.1)], series=[("year", "1"), (1850, 1.5)])
    events: list[ImportProgress] = []
    stats = run(Settings(db=clean_db.settings), tmp_path, progress=events.append)
    assert (stats.files, stats.title_updated, stats.crn_rows, stats.series_rows, stats.errors, stats.cancelled) == (1, 1, 1, 1, 0, False)
    assert clean_db.get_all_probes() == ["pp01"]
    assert clean_db.get_crn("pp01") == [CrnRow(1850, raw=1.1)]
    assert clean_db.get_series_matrix("pp01").series == ["1"]
    assert [e.stage for e in events] == ["scan", "file", "journal", "geometry"]
    assert events[2].count is None  # dendro_point_living.xlsx у папцы няма
    postgis = clean_db.has_postgis() and clean_db.has_geometry_column()
    assert (events[3].count is None) is (not postgis)
    assert clean_db.is_connected  # run() зачыняе толькі сваё злучэнне


def test_run_cancelled_before_first_file_skips_later_stages(clean_db, tmp_path):
    write_workbook(tmp_path / "PP01.xlsx", title=[("Широта,грд", 54.1)])
    events: list[ImportProgress] = []
    stats = run(Settings(db=clean_db.settings), tmp_path, progress=events.append, should_stop=lambda: True)
    assert stats.cancelled and stats.files == 0
    assert [e.stage for e in events] == ["scan"]
    assert clean_db.get_all_probes() == []
